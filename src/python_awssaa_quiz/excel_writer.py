"""理解度(◎○△)の判定と、Excelの「理解度」列への書き戻し。

書式・入力規則・「分野別集計」シートの数式を壊さないよう、pandasは使わず
openpyxlで「理解度」列のセルだけを書き換え、元ファイルとは別名で保存する。
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from datetime import date
from pathlib import Path

import openpyxl

from .loader import FIRST_DATA_ROW, SHEET_NAME, UNDERSTANDING_LABELS

NO_COLUMN = 1  # A列
UNDERSTANDING_COLUMN = 7  # G列


def judge_understanding(recent_results: Sequence[bool]) -> str | None:
    """回答履歴(新しい順)から理解度を判定する。

    直近3回すべて正解なら◎、直近3回に1回でも正解があれば○、それ以外は△。
    回答がなければNone(判定しない)。
    """
    last3 = list(recent_results[:3])
    if not last3:
        return None
    if len(last3) == 3 and all(last3):
        return "◎"
    if any(last3):
        return "○"
    return "△"


def updated_path(src: Path, today: date | None = None) -> Path:
    today = today or date.today()
    return src.with_name(f"{src.stem}_updated_{today:%Y%m%d}{src.suffix}")


def write_understanding(
    src: Path | str,
    levels: Mapping[int, str],
    dest: Path | str | None = None,
    today: date | None = None,
) -> tuple[Path, int]:
    """No → 理解度記号(◎/○/△)の対応を「理解度」列に書き込み、別名で保存する。

    levelsに含まれないNoのセルは変更しない。保存先のパスと更新したセル数を返す。
    """
    src = Path(src)
    dest = Path(dest) if dest else updated_path(src, today)
    if dest.resolve() == src.resolve():
        raise ValueError("元のExcelファイルは上書きしません")

    wb = openpyxl.load_workbook(src)  # 数式は文字列のまま保持される
    ws = wb[SHEET_NAME]
    updated = 0
    for row in range(FIRST_DATA_ROW, ws.max_row + 1):
        no = ws.cell(row=row, column=NO_COLUMN).value
        if no is None:
            continue
        mark = levels.get(int(no))
        if mark is None:
            continue
        ws.cell(row=row, column=UNDERSTANDING_COLUMN).value = UNDERSTANDING_LABELS[mark]
        updated += 1

    # openpyxlは数式の計算結果を保存しないので、開いたときに再計算させる
    wb.calculation.fullCalcOnLoad = True
    wb.save(dest)
    return dest, updated
