import hashlib
from datetime import date

import openpyxl
import pytest

from python_awssaa_quiz.excel_writer import judge_understanding, write_understanding
from python_awssaa_quiz.loader import load_keywords


@pytest.mark.parametrize(
    "recent, expected",
    [
        ([], None),
        ([True, True, True], "◎"),
        ([True, True, True, False], "◎"),  # 直近3回だけを見る
        ([True, True], "○"),  # 3回未満は◎にしない
        ([False, False, True], "○"),
        ([False], "△"),
        ([False, False, False, True], "△"),
    ],
)
def test_judge_understanding(recent, expected):
    assert judge_understanding(recent) == expected


def _sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


@pytest.fixture
def written(excel_path, tmp_path):
    before = _sha(excel_path)
    src = tmp_path / excel_path.name
    src.write_bytes(excel_path.read_bytes())
    levels = {1: "◎", 2: "○", 142: "△"}
    dest, count = write_understanding(src, levels, today=date(2026, 10, 14))
    assert _sha(excel_path) == before
    return src, dest, count


def test_saved_under_new_name_and_source_unchanged(written):
    src, dest, count = written
    assert dest.name == "AWS_SAA_重要キーワード集_updated_20261014.xlsx"
    assert count == 3
    assert load_keywords(src)[0].understanding is None  # 元ファイルはそのまま


def test_only_understanding_column_changes(written):
    src, dest, _ = written
    old_wb, new_wb = openpyxl.load_workbook(src), openpyxl.load_workbook(dest)
    for ws_old in old_wb.worksheets:
        ws_new = new_wb[ws_old.title]
        assert ws_new.max_row == ws_old.max_row and ws_new.max_column == ws_old.max_column
        for row_old, row_new in zip(ws_old.iter_rows(), ws_new.iter_rows()):
            for c_old, c_new in zip(row_old, row_new):
                if ws_old.title == "キーワード一覧" and c_old.column == 7 and c_old.row >= 5:
                    continue
                assert c_new.value == c_old.value, c_old.coordinate
                assert c_new.fill.fgColor.rgb == c_old.fill.fgColor.rgb
                assert c_new.font.b == c_old.font.b

    ws = new_wb["キーワード一覧"]
    assert ws["G5"].value == "◎ 説明できる"
    assert ws["G6"].value == "○ なんとなく"
    assert ws["G146"].value == "△ 要復習"  # No.142
    assert ws["G7"].value is None


def test_formulas_validation_and_formatting_kept(written):
    _, dest, _ = written
    wb = openpyxl.load_workbook(dest)
    summary = wb["分野別集計"]
    assert summary["B4"].value == "=COUNTIF(キーワード一覧!$B$5:$B$235,A4)"
    assert summary["D4"].value.startswith("=COUNTIFS(")
    assert summary["B16"].value == "=SUM(B4:B15)"
    formulas = [c.value for row in summary.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert len(formulas) == 13 * 6

    ws = wb["キーワード一覧"]
    dvs = ws.data_validations.dataValidation
    assert [str(dv.sqref) for dv in dvs] == ["G5:G235"]
    assert "◎ 説明できる" in dvs[0].formula1
    assert [str(cf.sqref) for cf in ws.conditional_formatting] == ["A5:G235"]


def test_reloaded_values_match_loader(written):
    _, dest, _ = written
    kws = load_keywords(dest)
    assert [k.understanding for k in kws[:3]] == ["◎", "○", None]
    assert kws[141].understanding == "△"


def test_refuses_to_overwrite_source(excel_path, tmp_path):
    src = tmp_path / "a.xlsx"
    src.write_bytes(excel_path.read_bytes())
    with pytest.raises(ValueError):
        write_understanding(src, {1: "◎"}, dest=src)
